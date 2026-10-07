"""Excel de « Ton entreprise » : modèle à remplir, export de la base, lecture d'un import.

Mêmes noms de feuilles et de colonnes que le kit Microsoft Lists (Listes-import.zip) : un dirigeant qui a déjà préparé
ses listes pour SharePoint peut importer les mêmes fichiers ici, et inversement l'export se ré-importe dans Microsoft Lists.
Fonctions pures (aucune base) : testées sans serveur.
"""
import io
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

# feuille → (colonnes, aide par colonne)
FEUILLES = {
    "Annuaire": (["Title", "Email", "Role", "Poste", "TelPro", "EmailPro"],
                 ["Prénom Nom", "E-mail de la personne (sert d'identifiant)", "Salarie ou Employeur", "Poste", "Téléphone pro", "E-mail pro"]),
    "TypesPieces": (["Title", "ParDefaut", "Consigne"],
                    ["Nom de la pièce (RIB, carte d'identité…)", "Oui = demandée à chaque nouvelle personne", "Indication pour la personne (facultatif)"]),
    "PiecesRequises": (["Title", "Salarie"],
                       ["Nom de la pièce", "Nom ou e-mail de la personne (comme dans Annuaire)"]),
    "Planning": (["Title", "Salarie", "DateEcheance", "Creneau", "Livrable", "Commentaire"],
                 ["Tâche ou événement", "Nom ou e-mail (vide = toute l'équipe)", "AAAA-MM-JJ", "Matin ou Après-midi (vide = journée)", "Livrable attendu", "Commentaire"]),
    "DossierSalarie": (["Title", "Salarie", "LienDrive"],
                       ["Nom", "Nom ou e-mail", "Lien https du dossier de la personne dans le Drive de l'entreprise"]),
}
EXEMPLES = {
    "Annuaire": ["Léa Martin", "lea.martin@exemple.fr", "Salarie", "Commerciale", "06 12 34 56 78", "lea.martin@exemple.fr"],
    "TypesPieces": ["RIB", "Oui", "Un RIB au nom de la personne"],
    "PiecesRequises": ["Carte vitale", "Léa Martin"],
    "Planning": ["Prospection", "Léa Martin", "2026-10-15", "Matin", "10 prospects qualifiés", ""],
    "DossierSalarie": ["Léa Martin", "Léa Martin", "https://exemple.sharepoint.com/sites/RH/Documents-RH/Lea-Martin"],
}
# Colonnes qui identifient une feuille venue du kit Microsoft (onglet « Liste »), sans la confondre avec Pieces, Absences…
REQUIS = {
    "Annuaire": {"Title", "Email", "Role"},
    "TypesPieces": {"Title", "ParDefaut"},
    "PiecesRequises": {"Title", "Salarie"},
    "Planning": {"Title", "Salarie", "DateEcheance"},
    "DossierSalarie": {"Salarie", "LienDrive"},
}
NOM_EXEMPLE = "Léa Martin"
EMAILS_EXEMPLE = {"lea.martin@exemple.fr", "lea.martin@entreprise.fr"}
_ENTETE = Font(bold=True, color="FFFFFF")
_FOND = PatternFill("solid", fgColor="1F3A5F")
MAX_LIGNES = 2000


def _entete(ws, colonnes):
    ws.append(colonnes)
    for c in ws[1]:
        c.font, c.fill = _ENTETE, _FOND
    for i, col in enumerate(colonnes):
        ws.column_dimensions[ws.cell(1, i + 1).column_letter].width = max(14, len(col) + 6)


def _octets(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def classeur_modele() -> bytes:
    wb = Workbook()
    lisez = wb.active
    lisez.title = "Lisez-moi"
    for ligne in (
        ["Ton entreprise by Zayado : modèle d'import"],
        [""],
        ["1. Remplis les onglets (une ligne d'exemple est fournie dans chacun : remplace-la ou supprime-la)."],
        ["2. Dans Zayado › Ton entreprise › Réglages › « Importer l'Excel », choisis ce fichier."],
        ["3. Les personnes de l'onglet Annuaire reçoivent une invitation (lien à envoyer depuis l'onglet Équipe)."],
        ["Rien n'est supprimé par un import : une personne déjà présente est mise à jour, jamais dupliquée."],
        ["Les fichiers (RIB, pièces d'identité…) restent dans TON Drive : l'Excel ne contient que des noms et des liens."],
    ):
        lisez.append(ligne)
    lisez.column_dimensions["A"].width = 110
    lisez["A1"].font = Font(bold=True, size=14)
    for nom, (colonnes, aides) in FEUILLES.items():
        ws = wb.create_sheet(nom)
        _entete(ws, colonnes)
        ws.append(EXEMPLES[nom])
        ws.append([])
        ws.append([f"Aide : {c} = {a}" for c, a in zip(colonnes, aides)])
        for c in ws[4]:
            c.font = Font(italic=True, color="808080")
    return _octets(wb)


def _txt(v) -> str:
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return v.date().isoformat() if isinstance(v, datetime) else v.isoformat()
    return str(v).strip()


def lire_classeur(octets: bytes) -> dict:
    """Lit les onglets connus. Accepte aussi un fichier d'une seule liste du kit Microsoft (onglet « Liste »),
    reconnu à ses colonnes. Lève ValueError avec un message lisible."""
    try:
        wb = load_workbook(io.BytesIO(octets), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001
        raise ValueError("Ce fichier n'est pas un Excel (.xlsx) lisible.")
    out = {}
    for ws in wb.worksheets:
        lignes = ws.iter_rows(values_only=True)
        try:
            entete = [_txt(c) for c in next(lignes)]
        except StopIteration:
            continue
        nom = ws.title if ws.title in FEUILLES else None
        if nom is None:  # fichier unitaire du kit (onglet « Liste ») : on le reconnaît à ses colonnes
            cols = {c for c in entete if c}
            for k, requis in REQUIS.items():
                if requis <= cols and (k != "PiecesRequises" or cols == requis):
                    nom = k
                    break
        if nom is None:
            continue
        donnees = []
        for n, ligne in enumerate(lignes):
            if n >= MAX_LIGNES:
                break
            d = {entete[i]: _txt(v) for i, v in enumerate(ligne) if i < len(entete) and entete[i]}
            if not any(d.values()) or any(str(v).startswith("Aide :") for v in d.values()):
                continue
            if d.get("Email", "").lower() in EMAILS_EXEMPLE or NOM_EXEMPLE in (d.get("Salarie"), d.get("Title") if nom == "Annuaire" else None):
                continue  # ligne d'exemple (Léa Martin) du modèle ou du kit Microsoft, oubliée dans le fichier
            donnees.append(d)
        out.setdefault(nom, []).extend(donnees)
    if not out:
        raise ValueError("Aucun onglet reconnu (Annuaire, TypesPieces, PiecesRequises, Planning, DossierSalarie).")
    return out


def classeur_export(org_nom, membres, noms, fiches, types, pieces, planning, absences, avec_perso: bool) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Annuaire"
    _entete(ws, ["Title", "Email", "Role", "Poste", "TelPro", "EmailPro", "Statut"])
    for c in membres:
        f = fiches.get(c.id)
        ws.append([noms.get(c.id, ""), c.email or "", "Employeur" if c.role in ("proprietaire", "manager") else "Salarie", c.poste or "",
                   (f.tel_pro if f else "") or "", (f.email_pro if f else "") or "", "Actif" if c.statut == "actif" else "Invité"])
    ws = wb.create_sheet("TypesPieces")
    _entete(ws, ["Title", "ParDefaut", "Consigne"])
    for t in types:
        ws.append([t.titre, "Oui" if t.par_defaut else "Non", t.consigne or ""])
    ws = wb.create_sheet("Pieces")
    _entete(ws, ["Title", "Salarie", "Statut", "LienFichier", "Commentaire"])
    lib = {"a_fournir": "À fournir", "fournie": "En attente", "validee": "Validé", "refusee": "Refusé"}
    for p in pieces:
        ws.append([p.titre, noms.get(p.membre_id, ""), lib.get(p.statut, p.statut), p.reference or "", p.commentaire or ""])
    ws = wb.create_sheet("Planning")
    _entete(ws, ["Title", "Salarie", "DateEcheance", "Fin", "Creneau", "Commentaire", "Statut"])
    cren = {"matin": "Matin", "apres_midi": "Après-midi"}
    st = {"propose": "Proposé", "vu": "Vu"}
    for r in sorted(planning, key=lambda x: x.debut):
        ws.append([r.titre, noms.get(r.membre_id, "Toute l'équipe") if r.membre_id else "Toute l'équipe", r.debut, r.fin or "",
                   cren.get(r.creneau or "", "Journée"), r.note or "", st.get(r.statut or "", "Posé")])
    ws = wb.create_sheet("Absences")
    _entete(ws, ["Title", "Salarie", "Du", "Au", "Statut"])
    sa = {"demandee": "En attente", "acceptee": "Acceptée", "refusee": "Refusée"}
    for a in absences:
        ws.append([a.type, noms.get(a.membre_id, ""), a.debut, a.fin, sa.get(a.statut, a.statut)])
    ws = wb.create_sheet("DossierSalarie")
    colonnes = ["Title", "Salarie", "LienDrive"] + (["Adresse", "CodePostal", "Ville", "TelPerso", "EmailPerso", "UrgenceNom", "UrgenceTel"] if avec_perso else [])
    _entete(ws, colonnes)
    for c in membres:
        f = fiches.get(c.id)
        ligne = [noms.get(c.id, ""), noms.get(c.id, ""), (f.lien_drive if f else "") or ""]
        if avec_perso:
            ligne += [(getattr(f, k) if f else "") or "" for k in ("adresse", "code_postal", "ville", "tel_perso", "email_perso", "urgence_nom", "urgence_tel")]
        ws.append(ligne)
    return _octets(wb)
