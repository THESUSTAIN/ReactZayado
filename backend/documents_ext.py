"""Documents créés par l'IA : Word, Excel, Markdown, CSV, texte, et images.

- Dossier de rangement : l'utilisateur colle l'URL d'un dossier de son Google Drive ou de
  son OneDrive / SharePoint (demandé la première fois qu'il enregistre). On le mémorise.
- « Mes documents » : ouvre ce dossier (ou la racine du Drive relié).
- Générer : à partir du texte de l'IA (Markdown), on fabrique le fichier dans le format choisi,
  puis on le télécharge OU on le range dans le dossier du Drive.
- Images : génération par le modèle d'images (Mammouth), téléchargement ou rangement dans le Drive.
"""
import base64
import csv
import io
import json
import logging
import re
from typing import Optional
from urllib.parse import quote, urlparse

import httpx
from fastapi import Depends, HTTPException
from fastapi.responses import Response
from pydantic import BaseModel, Field
from sqlalchemy.ext.asyncio import AsyncSession

log = logging.getLogger("kairos.documents")

FORMATS = {
    "md": ("text/markdown", ".md"),
    "txt": ("text/plain", ".txt"),
    "csv": ("text/csv", ".csv"),
    "docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", ".docx"),
    "xlsx": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", ".xlsx"),
}


def nom_fichier(titre: str, ext: str) -> str:
    base = re.sub(r"[^\w.\- ]+", "", (titre or "document"), flags=re.UNICODE).strip().replace(" ", "-")[:90] or "document"
    return base + ext


def _tableaux_markdown(md: str) -> list[list[list[str]]]:
    """Extrait les tableaux Markdown (| a | b |) en listes de lignes."""
    tableaux, courant = [], []
    for ligne in md.splitlines():
        l = ligne.strip()
        if l.startswith("|") and l.endswith("|") and len(l) > 2:
            cellules = [c.strip() for c in l.strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cellules if c):
                continue  # ligne de séparation
            courant.append(cellules)
        elif courant:
            tableaux.append(courant)
            courant = []
    if courant:
        tableaux.append(courant)
    return tableaux


def _nombre(v: str):
    if "%" in v:
        return v
    t = v.replace(" ", "").replace(" ", "").replace("€", "").replace("%", "")
    t = t.replace(",", ".") if re.fullmatch(r"-?\d+(,\d+)?", t) else t
    try:
        return float(t) if "." in t else int(t)
    except ValueError:
        return v


def _sans_markdown(t: str) -> str:
    return re.sub(r"(\*\*|__|\*|_|`)", "", t)


def fabriquer(contenu: str, fmt: str, titre: str) -> bytes:
    if fmt in ("md", "txt"):
        return (contenu if fmt == "md" else _sans_markdown(contenu)).encode("utf-8")
    if fmt == "csv":
        tab = (_tableaux_markdown(contenu) or [[[l] for l in contenu.splitlines() if l.strip()]])[0]
        out = io.StringIO()
        csv.writer(out, delimiter=";").writerows(tab)
        return ("﻿" + out.getvalue()).encode("utf-8")
    if fmt == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook()
        tableaux = _tableaux_markdown(contenu)
        if not tableaux:
            tableaux = [[["Contenu"]] + [[_sans_markdown(l)] for l in contenu.splitlines() if l.strip()]]
        for i, tab in enumerate(tableaux):
            ws = wb.active if i == 0 else wb.create_sheet()
            ws.title = (f"Tableau {i + 1}" if len(tableaux) > 1 else (titre or "Feuille"))[:31].replace("/", "-")
            for r, ligne in enumerate(tab, start=1):
                for c, val in enumerate(ligne, start=1):
                    cel = ws.cell(row=r, column=c, value=_sans_markdown(val) if r == 1 else _nombre(_sans_markdown(val)))
                    if r == 1:
                        cel.font = Font(bold=True, color="FFFFFF")
                        cel.fill = PatternFill("solid", fgColor="1D3A70")
            for col in ws.columns:
                largeur = max(len(str(c.value or "")) for c in col)
                ws.column_dimensions[col[0].column_letter].width = min(60, max(10, largeur + 2))
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()
    if fmt == "docx":
        from docx import Document
        doc = Document()
        premier = next((l.strip() for l in contenu.splitlines() if l.strip()), "")
        if titre and not re.match(r"^#\s+" + re.escape(titre.strip()) + r"\s*$", premier):
            doc.add_heading(titre, level=0)
        lignes = contenu.splitlines()
        i = 0
        while i < len(lignes):
            l = lignes[i].rstrip()
            s = l.strip()
            if s.startswith("|") and s.endswith("|"):
                bloc = []
                while i < len(lignes) and lignes[i].strip().startswith("|"):
                    bloc.append(lignes[i])
                    i += 1
                tab = _tableaux_markdown("\n".join(bloc))
                if tab:
                    t = tab[0]
                    table = doc.add_table(rows=len(t), cols=max(len(x) for x in t))
                    table.style = "Table Grid"
                    for r, ligne in enumerate(t):
                        for c, val in enumerate(ligne):
                            table.cell(r, c).text = _sans_markdown(val)
                continue
            m = re.match(r"^(#{1,4})\s+(.*)", s)
            if m:
                doc.add_heading(_sans_markdown(m.group(2)), level=min(len(m.group(1)), 4))
            elif re.match(r"^[-*•]\s+", s):
                doc.add_paragraph(_sans_markdown(re.sub(r"^[-*•]\s+", "", s)), style="List Bullet")
            elif re.match(r"^\d+[.)]\s+", s):
                doc.add_paragraph(_sans_markdown(re.sub(r"^\d+[.)]\s+", "", s)), style="List Number")
            elif s:
                p = doc.add_paragraph()
                for morceau in re.split(r"(\*\*[^*]+\*\*)", s):
                    if morceau.startswith("**") and morceau.endswith("**"):
                        p.add_run(morceau[2:-2]).bold = True
                    else:
                        p.add_run(_sans_markdown(morceau))
            i += 1
        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()
    raise HTTPException(422, "Format non pris en charge.")


def analyser_dossier(url: str) -> dict:
    """Reconnaît un dossier Google Drive ou OneDrive / SharePoint à partir de son URL."""
    url = (url or "").strip()
    u = urlparse(url)
    if u.scheme != "https" or not u.hostname:
        raise HTTPException(422, "Colle l'adresse complète du dossier (elle commence par https://).")
    hote = u.hostname.lower()
    if hote == "drive.google.com":
        m = re.search(r"/folders/([\w-]{10,})", u.path)
        if not m:
            raise HTTPException(422, "Ce lien Google Drive n'est pas un dossier : ouvre le dossier puis copie l'adresse.")
        return {"provider": "google", "url": url, "id": m.group(1)}
    if hote.endswith("sharepoint.com") or hote in ("onedrive.live.com", "1drv.ms") or hote.endswith(".onedrive.com"):
        return {"provider": "microsoft", "url": url, "id": None}
    raise HTTPException(422, "Je reconnais les dossiers Google Drive, OneDrive et SharePoint.")


def install_documents(g: dict) -> None:
    api, get_db, _uid = g["api"], g["get_db"], g["_uid"]

    async def _cm(db) -> dict:
        p = await g["_profil"](db, _uid())
        return dict(p.contexte_metier or {}), p

    async def _connecte(db, provider: str) -> bool:
        try:
            await g["_cloud_token"](db, provider)
            return True
        except HTTPException:
            return False

    @api.get("/documents/dossier")
    async def lire_dossier(db: AsyncSession = Depends(get_db)):
        cm, _ = await _cm(db)
        d = cm.get("documents_dossier")
        return {"dossier": d, "dossier_auto": cm.get("documents_dossier_auto"), "google": await _connecte(db, "google"), "microsoft": await _connecte(db, "microsoft"),
                "formats": list(FORMATS), "images": bool(g.get("images_disponibles") and g["images_disponibles"]())}

    class DossierIn(BaseModel):
        url: Optional[str] = Field(default=None, max_length=1000)

    @api.put("/documents/dossier")
    async def regler_dossier(body: DossierIn, db: AsyncSession = Depends(get_db)):
        cm, p = await _cm(db)
        if not body.url:
            cm.pop("documents_dossier", None)
        else:
            d = analyser_dossier(body.url)
            if d["provider"] == "microsoft" and await _connecte(db, "microsoft"):
                # Lien de partage → dossier réel (drive + élément) via Microsoft Graph.
                token, _ = await g["_cloud_token"](db, "microsoft")
                code = "u!" + base64.urlsafe_b64encode(d["url"].encode()).decode().rstrip("=")
                async with httpx.AsyncClient(timeout=20) as c:
                    r = await c.get(f"https://graph.microsoft.com/v1.0/shares/{code}/driveItem", headers={"Authorization": f"Bearer {token}"})
                if r.status_code < 400:
                    it = r.json()
                    d["id"], d["drive_id"] = it.get("id"), (it.get("parentReference") or {}).get("driveId")
                    d["nom"] = it.get("name")
                else:
                    raise HTTPException(422, "Je n'arrive pas à ouvrir ce dossier OneDrive / SharePoint avec ton compte relié.")
            cm["documents_dossier"] = d
            cm["document_provider"] = d["provider"]
        p.contexte_metier = cm
        await db.commit()
        return {"dossier": cm.get("documents_dossier")}

    @api.get("/documents/ouvrir")
    async def ouvrir(db: AsyncSession = Depends(get_db)):
        cm, _ = await _cm(db)
        d = cm.get("documents_dossier") or cm.get("documents_dossier_auto")
        if d and d.get("url"):
            return {"url": d["url"]}
        provider = cm.get("document_provider") or ("google" if await _connecte(db, "google") else "microsoft" if await _connecte(db, "microsoft") else None)
        if provider == "google":
            return {"url": "https://drive.google.com/drive/my-drive"}
        if provider == "microsoft":
            return {"url": "https://www.office.com/launch/onedrive"}
        return {"url": None}

    async def _dossier_zayado_google(db, c, token: str, cm: dict) -> str:
        """Dossier « Zayado » créé par l'app dans le Drive (accès garanti avec le droit drive.file)."""
        auto = cm.get("documents_dossier_auto") or {}
        if auto.get("provider") == "google" and auto.get("id"):
            return auto["id"]
        r = await c.post("https://www.googleapis.com/drive/v3/files?fields=id,webViewLink",
                         headers={"Authorization": f"Bearer {token}"},
                         json={"name": "Zayado", "mimeType": "application/vnd.google-apps.folder"})
        if r.status_code >= 400:
            raise HTTPException(502, "Impossible de créer le dossier Zayado dans ton Drive.")
        j = r.json()
        cm["documents_dossier_auto"] = {"provider": "google", "id": j["id"], "url": j.get("webViewLink")}
        p = await g["_profil"](db, _uid())
        p.contexte_metier = {**(p.contexte_metier or {}), "documents_dossier_auto": cm["documents_dossier_auto"]}
        await db.commit()
        return j["id"]

    async def _ranger(db, cm: dict, nom: str, octets: bytes, mime: str) -> dict:
        d = cm.get("documents_dossier") or {}
        provider = d.get("provider") or cm.get("document_provider") or "google"
        try:
            token, _ = await g["_cloud_token"](db, provider)
        except HTTPException:
            raise HTTPException(409, "Relie d'abord ton Google Drive ou ton OneDrive (Paramètres › Connexions).")
        async with httpx.AsyncClient(timeout=60) as c:
            if provider == "google":
                meta = {"name": nom}
                if d.get("id") and d.get("provider") == "google":
                    meta["parents"] = [d["id"]]
                else:
                    meta["parents"] = [await _dossier_zayado_google(db, c, token, cm)]
                r = await c.post("https://www.googleapis.com/upload/drive/v3/files?uploadType=multipart&fields=id,name,webViewLink&supportsAllDrives=true",
                                 headers={"Authorization": f"Bearer {token}"},
                                 files={"metadata": ("metadata.json", json.dumps(meta), "application/json"), "file": (nom, octets, mime)})
            else:
                if d.get("id") and d.get("drive_id"):
                    cible = f"https://graph.microsoft.com/v1.0/drives/{d['drive_id']}/items/{d['id']}:/{quote(nom)}:/content"
                else:
                    cible = f"https://graph.microsoft.com/v1.0/me/drive/root:/{quote('Zayado/' + nom, safe='/')}:/content"
                r = await c.put(cible, headers={"Authorization": f"Bearer {token}", "Content-Type": mime}, content=octets)
        if r.status_code >= 400:
            log.warning("Rangement Drive refusé (%s) : %s", r.status_code, r.text[:300])
            if r.status_code in (403, 404) and d.get("id"):
                raise HTTPException(409, "Zayado n'a pas accès à ce dossier. Choisis un dossier créé par Zayado ou partage-le avec ton compte relié.")
            raise HTTPException(502, "Le Drive a refusé l'enregistrement. Réessaie dans un instant.")
        j = r.json()
        return {"ok": True, "provider": provider, "nom": nom, "url": j.get("webViewLink") or j.get("webUrl")}

    async def ranger_document(db, titre: str, contenu: str, fmt: str = "docx") -> dict:
        """Fabrique le document (Word par défaut) et le range dans le Drive / OneDrive relié de l'utilisateur courant.
        Utilisé par les agents : dans l'appli, on ne télécharge pas, le livrable va directement dans le Drive."""
        mime, ext = FORMATS[fmt]
        cm, _ = await _cm(db)
        return await _ranger(db, cm, nom_fichier(titre, ext), fabriquer(contenu, fmt, titre.strip()), mime)

    g["ranger_document_drive"] = ranger_document

    class DocIn(BaseModel):
        titre: str = Field(default="Document", max_length=150)
        contenu: str = Field(min_length=1, max_length=200000)
        format: str = "docx"
        destination: str = "telecharger"  # telecharger | drive

    @api.post("/documents/generer")
    async def generer(body: DocIn, db: AsyncSession = Depends(get_db)):
        if body.format not in FORMATS:
            raise HTTPException(422, "Format non pris en charge (Word, Excel, Markdown, CSV, texte).")
        mime, ext = FORMATS[body.format]
        octets = fabriquer(body.contenu, body.format, body.titre.strip())
        nom = nom_fichier(body.titre, ext)
        if body.destination == "drive":
            cm, _ = await _cm(db)
            return await _ranger(db, cm, nom, octets, mime)
        return Response(content=octets, media_type=mime,
                        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(nom)}"})

    class ImageIn(BaseModel):
        description: str = Field(min_length=3, max_length=2000)
        destination: str = "telecharger"

    @api.post("/documents/image")
    async def image(body: ImageIn, db: AsyncSession = Depends(get_db)):
        f = g.get("_generer_image")
        if not f:
            raise HTTPException(503, "Les images IA ne sont pas disponibles sur ce serveur.")
        octets, mime = await f(body.description.strip(), libre=True)
        ext = {"image/jpeg": ".jpg", "image/webp": ".webp"}.get(mime, ".png")
        nom = nom_fichier(body.description[:60], ext)
        if body.destination == "drive":
            cm, _ = await _cm(db)
            return await _ranger(db, cm, nom, octets, mime)
        return {"ok": True, "nom": nom, "mime": mime, "data": base64.b64encode(octets).decode()}

    class FichierIn(BaseModel):
        nom: str = Field(min_length=1, max_length=150)
        mime: str = Field(max_length=100)
        data: str = Field(max_length=15_000_000)  # base64

    @api.post("/documents/ranger-fichier")
    async def ranger_fichier(body: FichierIn, db: AsyncSession = Depends(get_db)):
        if not re.fullmatch(r"(image/(png|jpeg|webp)|text/(markdown|plain|csv)|application/vnd\.openxmlformats-officedocument\.[\w.]+)", body.mime):
            raise HTTPException(422, "Type de fichier non pris en charge.")
        try:
            octets = base64.b64decode(body.data, validate=True)
        except Exception:  # noqa: BLE001
            raise HTTPException(422, "Fichier illisible.")
        cm, _ = await _cm(db)
        return await _ranger(db, cm, re.sub(r"[^\w.\- ]+", "", body.nom)[:150] or "fichier", octets, body.mime)
