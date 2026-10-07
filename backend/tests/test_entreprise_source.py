"""« J'ai déjà mes fichiers » : le classeur du client, la correspondance proposée, la synchronisation."""
import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import server  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from entreprise_source_ext import lire_xlsx, proposer, convertir  # noqa: E402
from tests.test_entreprise import _equipe  # noqa: E402


def classeur_client() -> bytes:
    """Un classeur « à la façon du client » : ses noms d'onglets et de colonnes, pas ceux de Zayado."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Salariés"
    ws.append(["Nom Prénom", "Adresse mail", "Fonction", "Portable"])
    ws.append(["Durand Paul", "paul@client.fr", "Comptable", "0611"])
    ws.append(["Roux Inès", "ines@client.fr", "Vendeuse", "0622"])
    p = wb.create_sheet("Agenda")
    p.append(["Activité", "Jour", "Collaborateur", "Détail"])
    p.append(["Inventaire", "2026-11-03", "Durand Paul", "rayon frais"])
    d = wb.create_sheet("Documents")
    d.append(["Pièce", "Obligatoire"])
    d.append(["RIB", "Oui"])
    d.append(["Permis", "Non"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_proposition_sur_les_noms_du_client():
    o = lire_xlsx(classeur_client())
    p = proposer(o)
    assert p["personnes"] == {"onglet": "Salariés", "colonnes": {"nom": "Nom Prénom", "email": "Adresse mail", "poste": "Fonction", "tel_pro": "Portable"}}
    assert p["planning"]["onglet"] == "Agenda" and p["planning"]["colonnes"]["date"] == "Jour" and p["planning"]["colonnes"]["personne"] == "Collaborateur"
    assert p["pieces"]["onglet"] == "Documents" and p["pieces"]["colonnes"]["titre"] == "Pièce"
    d = convertir(o, p)
    assert d["Annuaire"][0]["Email"] == "paul@client.fr" and d["TypesPieces"][1]["ParDefaut"] == "Non" and d["Planning"][0]["DateEcheance"] == "2026-11-03"


def test_relier_puis_synchroniser(client, compte, monkeypatch):
    a = _equipe(client, compte)

    async def chercher(fournisseur, q):
        return [{"id": "F1", "nom": "RH entreprise.xlsx", "url": "https://onedrive/F1"}]

    async def telecharger(fournisseur, fid):
        assert fid == "F1"
        return classeur_client()
    monkeypatch.setitem(server.__dict__, "_source_test_chercher", chercher)
    monkeypatch.setitem(server.__dict__, "_source_test_telecharger", telecharger)
    assert client.get("/api/entreprise/source/fichiers?fournisseur=microsoft", headers=a["manager"]).status_code == 403
    f = client.get("/api/entreprise/source/fichiers?fournisseur=microsoft", headers=a["proprietaire"]).json()["fichiers"][0]
    an = client.post("/api/entreprise/source/analyser", json={"fournisseur": "microsoft", "fichier_id": f["id"]}, headers=a["proprietaire"]).json()
    assert an["onglets"]["Salariés"]["exemples"][0][0] == "Durand Paul" and an["proposition"]["personnes"]["colonnes"]["email"] == "Adresse mail"
    # une correspondance incomplète est refusée (champ obligatoire manquant)
    mauvais = {"personnes": {"onglet": "Salariés", "colonnes": {"nom": "Nom Prénom"}}}
    assert client.put("/api/entreprise/source", json={"fournisseur": "microsoft", "fichier": f, "correspondance": mauvais}, headers=a["proprietaire"]).status_code == 422
    assert client.put("/api/entreprise/source", json={"fournisseur": "microsoft", "fichier": f, "correspondance": an["proposition"]}, headers=a["proprietaire"]).json()["ok"]
    r = client.post("/api/entreprise/source/synchroniser", headers=a["proprietaire"]).json()
    assert r["invites"] == 2 and r["types"] == 2 and r["planning"] == 1
    # relancer : personne n'est dupliqué
    r2 = client.post("/api/entreprise/source/synchroniser", headers=a["proprietaire"]).json()
    assert r2["invites"] == 0 and r2["mis_a_jour"] == 2
    assert client.get("/api/entreprise/source", headers=a["proprietaire"]).json()["source"]["derniere_synchro"]
