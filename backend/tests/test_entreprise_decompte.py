"""Décompte du mois : absences (dont CP), temps école, paie estimée (total / déjà versé / reste), droits de lecture, export."""
import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import server  # noqa: E402,F401
from entreprise_decompte import calculer, jours_ouvres, bornes_mois  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from tests.test_entreprise import _equipe, _ids  # noqa: E402

REMU = {"mode": "forfait_jour", "taux": 100, "heures_jour": 7, "cp_initial": 5, "cp_depuis": "2026-09-01", "ecole_jours": [2], "ecole_payee": False}


def test_calcul_pur_octobre_2026():
    assert len(jours_ouvres(*bornes_mois("2026-10"))) == 22
    abs_ = [{"type": "conges", "debut": "2026-10-05", "fin": "2026-10-06"}, {"type": "maladie", "debut": "2026-10-15", "fin": "2026-10-15"}]
    d = calculer("2026-10", [], abs_, [], [], REMU, 2.08, [a for a in abs_ if a["type"] == "conges"], 500, True)
    # 22 jours ouvrés ; mercredis d'école : 7, 14, 21, 28 octobre = 4 ; CP 2 ; maladie 1
    assert d["absences"]["ecole"] == 4 and d["absences"]["conges"] == 2 and d["absences"]["maladie"] == 1
    assert d["jours_travailles"] == 22 - 7
    # CP : 5 de départ + 2 mois × 2,08 = 9,16 acquis ; 2 pris ; solde 7,16
    assert d["cp"] == {"acquis": 9.16, "pris": 2, "pris_ce_mois": 2, "solde": 7.16}
    # Payé : 15 travaillés + 2 CP (école non payée, maladie non payée) = 17 jours × 100 €
    p = d["paie"]
    assert p["quantite"] == 17 and p["total_a_payer"] == 1700 and p["deja_verse"] == 500 and p["reste_a_verser"] == 1200
    # Alternant : l'école est payée
    d2 = calculer("2026-10", [], abs_, [], [], {**REMU, "ecole_payee": True}, 2.08, [], 0, True)
    assert d2["paie"]["quantite"] == 21
    # Sans droit de voir la paie : pas de montant
    assert calculer("2026-10", [], abs_, [], [], REMU, 2.08, [], 0, False)["paie"] is None


def test_decompte_droits_et_export(client, compte):
    a = _equipe(client, compte)
    ids = _ids(client, a["proprietaire"])
    assert client.put(f"/api/entreprise/remunerations/{ids['membre']}", json=REMU, headers=a["manager"]).status_code == 403
    assert client.put(f"/api/entreprise/remunerations/{ids['membre']}", json=REMU, headers=a["proprietaire"]).status_code == 200
    client.post("/api/entreprise/absences", json={"type": "conges", "debut": "2026-10-05", "fin": "2026-10-06", "membre_id": ids["membre"]}, headers=a["proprietaire"])
    client.put("/api/entreprise/versements", json={"membre_id": ids["membre"], "mois": "2026-10", "montant": 300}, headers=a["proprietaire"])
    # la personne voit SON décompte, paie comprise
    d = client.get("/api/entreprise/decompte?mois=2026-10", headers=a["membre"]).json()["decomptes"]
    assert len(d) == 1 and d[0]["absences"]["conges"] == 2 and d[0]["cp"]["pris_ce_mois"] == 2
    assert d[0]["paie"]["deja_verse"] == 300 and d[0]["paie"]["reste_a_verser"] == d[0]["paie"]["total_a_payer"] - 300
    assert client.get(f"/api/entreprise/remunerations/moi", headers=a["membre"]).json()["ecole_jours"] == [2]
    # pas celui des autres
    assert client.get(f"/api/entreprise/decompte?mois=2026-10&membre_id={ids['manager']}", headers=a["membre"]).status_code == 403
    assert client.get("/api/entreprise/decompte?mois=2026-10", headers=a["partenaire"]).status_code == 403
    # le manager voit l'équipe, sans les montants ; le propriétaire voit tout
    eq = {x["membre_id"]: x for x in client.get("/api/entreprise/decompte?mois=2026-10", headers=a["manager"]).json()["decomptes"]}
    assert ids["membre"] in eq and eq[ids["membre"]]["paie"] is None
    assert client.get(f"/api/entreprise/remunerations/{ids['membre']}", headers=a["manager"]).status_code == 403
    eq = {x["membre_id"]: x for x in client.get("/api/entreprise/decompte?mois=2026-10", headers=a["proprietaire"]).json()["decomptes"]}
    assert eq[ids["membre"]]["paie"]["total_a_payer"] > 0
    # le dirigeant aussi a son propre décompte
    assert len(client.get("/api/entreprise/decompte?mois=2026-10&membre_id=moi", headers=a["proprietaire"]).json()["decomptes"]) == 1
    # export Excel : un onglet par personne + synthèse
    r = client.get("/api/entreprise/decompte/export?mois=2026-10", headers=a["proprietaire"])
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames[0] == "Synthèse" and len(wb.sheetnames) >= 4
    r = client.get("/api/entreprise/decompte/export?mois=2026-10", headers=a["membre"])
    assert r.status_code == 200 and len(load_workbook(io.BytesIO(r.content)).sheetnames) == 2
    assert client.get("/api/entreprise/decompte?mois=octobre", headers=a["membre"]).status_code == 422
    # réglage CP de l'entreprise
    assert client.patch("/api/entreprise/parametres", json={"cp_par_mois": 2.5}, headers=a["proprietaire"]).status_code == 200
    d = client.get("/api/entreprise/decompte?mois=2026-10", headers=a["membre"]).json()["decomptes"][0]
    assert d["cp"]["acquis"] == 5 + 2 * 2.5
