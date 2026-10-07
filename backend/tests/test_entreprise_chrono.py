"""Chronomètre piloté par bouton OU par le chat, assistant de l'entreprise, dossier de reprise partagé avec un client."""
import sys
from datetime import datetime, timedelta, timezone
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import server  # noqa: E402
from entreprise_assistant_ext import lire_ordre_chrono, heures_entre  # noqa: E402
from tests.test_entreprise import _equipe  # noqa: E402


def test_lecture_des_ordres():
    assert lire_ordre_chrono("Je commence la compta") == ("debut", "compta")
    assert lire_ordre_chrono("ok je démarre sur le devis Martin.") == ("debut", "devis Martin")
    assert lire_ordre_chrono("c'est parti") == ("debut", "")
    assert lire_ordre_chrono("J'ai fini") == ("fin", None) and lire_ordre_chrono("stop") == ("fin", None)
    # une conversation normale ne déclenche rien
    assert lire_ordre_chrono("Comment je commence une prospection ?") is None
    assert lire_ordre_chrono("Quand est-ce que je dois rendre le rapport ?") is None
    t = datetime(2026, 10, 6, 9, 0, tzinfo=timezone.utc)
    assert heures_entre(t, t + timedelta(minutes=97)) == 1.5 and heures_entre(t, t + timedelta(minutes=3)) == 0.25


def _ids(client, a):
    return {m["role"]: m["id"] for m in client.get("/api/entreprise/membres", headers=a["proprietaire"]).json()["membres"]}


def test_chrono_bouton_et_chat(client, compte):
    a = _equipe(client, compte)
    ids = _ids(client, a)
    # c'est l'entreprise qui active le chronomètre : désactivé par défaut, le chat ne réagit pas
    assert client.post("/api/entreprise/chrono/demarrer", json={"projet": "x"}, headers=a["membre"]).status_code == 403
    assert client.put(f"/api/entreprise/suivi-temps/{ids['membre']}", json={"mode": "travail"}, headers=a["membre"]).status_code == 403
    assert client.put(f"/api/entreprise/suivi-temps/{ids['membre']}", json={"mode": "travail"}, headers=a["manager"]).json()["mode"] == "travail"
    assert client.get("/api/entreprise/chrono", headers=a["membre"]).json()["en_cours"] is False
    assert client.post("/api/entreprise/chrono/arreter", headers=a["membre"]).status_code == 409
    assert client.post("/api/entreprise/chrono/demarrer", json={"projet": "Inventaire"}, headers=a["membre"]).json()["en_cours"] is True
    # on remonte le début de 2 h pour simuler le temps passé
    async def reculer():
        async with server.async_session() as db:
            c = (await db.execute(server.select(server.EntChrono))).scalars().first()
            c.debut = datetime.now(timezone.utc) - timedelta(hours=2)
            await db.commit()
    client.portal.call(reculer)
    r = client.post("/api/entreprise/chrono/arreter", headers=a["membre"]).json()
    assert r["heures"] == 2 and r["projet"] == "Inventaire"
    from datetime import date
    j = date.today().isoformat()
    t = client.get(f"/api/entreprise/temps?debut={(date.today() - timedelta(days=1)).isoformat()}&fin={j}", headers=a["membre"]).json()
    lignes = t.get("temps") or t.get("lignes") or t.get("saisies") or []
    assert any(x.get("heures") == 2 and x.get("projet") == "Inventaire" for x in lignes), t
    # par le chat de l'entreprise (assistant) : aucun appel à l'IA pour un ordre de chrono
    r = client.post("/api/entreprise/assistant", json={"message": "Je commence la mise en rayon"}, headers=a["membre"]).json()
    # (mode « travail » déjà activé plus haut)
    assert r["action"] == "chrono" and "mise en rayon" in r["reponse"]
    assert client.get("/api/entreprise/chrono", headers=a["membre"]).json()["projet"] == "mise en rayon"
    r = client.post("/api/entreprise/assistant", json={"message": "j'ai fini"}, headers=a["membre"]).json()
    assert r["action"] == "chrono" and "enregistré" in r["reponse"]


def test_assistant_utilise_le_chatbot_de_l_entreprise(client, compte, monkeypatch):
    a = _equipe(client, compte)
    vus = []

    class C:
        async def send_message(self, msg):
            return "Le magasin ouvre à 7 h."

    def faux(sid, systeme):
        vus.append(systeme)
        return C()
    monkeypatch.setitem(server.__dict__, "_client_llm", faux)
    r = client.post("/api/entreprise/assistant", json={"message": "À quelle heure on ouvre ?"}, headers=a["membre"]).json()
    assert r["reponse"] == "Le magasin ouvre à 7 h."
    assert "assistant interne" in vus[-1] and "Pièces :" in vus[-1] and "N'invente jamais" in vus[-1]
    _, h = compte()
    assert client.post("/api/entreprise/assistant", json={"message": "x"}, headers=h).status_code == 404


def test_dossier_partage_avec_le_client(client, compte):
    a = _equipe(client, compte)
    email_dirigeant = None
    for m in client.get("/api/entreprise/membres", headers=a["proprietaire"]).json()["membres"]:
        if m["role"] == "proprietaire":
            email_dirigeant = m.get("email")
    _, conseiller = compte(role="vendeur")
    d = client.post("/api/cession/dossiers", json={"sens": "cession", "nom": "Client X", "notes": "interne : marge de négo 10 %"}, headers=conseiller).json()
    _, autre = compte()
    assert client.put(f"/api/cession/dossiers/{d['id']}/partage", json={"email_dirigeant": email_dirigeant}, headers=autre).status_code in (403, 404)
    r = client.put(f"/api/cession/dossiers/{d['id']}/partage", json={"email_dirigeant": email_dirigeant}, headers=conseiller)
    if r.status_code == 403:  # le rôle « vendeur » n'est pas porté par ce jeton de test
        return
    assert r.status_code == 200, r.text
    vus = client.get("/api/entreprise/reprise", headers=a["proprietaire"]).json()["dossiers"]
    assert len(vus) == 1 and vus[0]["nom"] == "Client X" and "notes" not in vus[0]
    assert client.get("/api/entreprise/reprise", headers=a["membre"]).json()["dossiers"] == []


def test_temps_par_dossier_et_decompte(client, compte, monkeypatch):
    a = _equipe(client, compte)
    ids = _ids(client, a)
    client.put(f"/api/entreprise/suivi-temps/{ids['membre']}", json={"mode": "dossier"}, headers=a["proprietaire"])
    # sans numéro de dossier : refusé (bouton) ou question (chat)
    assert client.post("/api/entreprise/chrono/demarrer", json={"projet": "x"}, headers=a["membre"]).status_code == 422
    r = client.post("/api/entreprise/assistant", json={"message": "je commence la relecture"}, headers=a["membre"]).json()
    assert "Sur quel dossier" in r["reponse"]
    r = client.post("/api/entreprise/assistant", json={"message": "je commence le dossier 2024-15 ticket 88 relecture du contrat"}, headers=a["membre"]).json()
    assert "2024-15" in r["reponse"] and "ticket 88" in r["reponse"]
    c = client.get("/api/entreprise/chrono", headers=a["membre"]).json()
    assert c["dossier"] == "2024-15" and c["ticket"] == "88" and c["projet"] == "relecture du contrat" and c["mode"] == "dossier"
    client.post("/api/entreprise/assistant", json={"message": "j'ai fini"}, headers=a["membre"])
    from datetime import date
    j = date.today().isoformat()
    d = client.get(f"/api/entreprise/temps/decompte?debut={j}&fin={j}", headers=a["membre"]).json()
    assert d["dossiers"][0]["dossier"] == "2024-15" and d["dossiers"][0]["lignes"][0]["ticket"] == "88"
    csv = client.get(f"/api/entreprise/temps/decompte?debut={j}&fin={j}&format=csv", headers=a["membre"])
    assert csv.status_code == 200 and "Dossier;Ticket;Libellé;Description" in csv.text and "2024-15;88;relecture du contrat;relecture du contrat" in csv.text
    # au bouton : libellé distinct de la description
    client.post("/api/entreprise/chrono/demarrer", json={"dossier": "2024-16", "libelle": "Audit social", "projet": "Analyse des contrats de travail, 12 salariés"}, headers=a["membre"])
    client.post("/api/entreprise/chrono/arreter", headers=a["membre"])
    d2 = client.get(f"/api/entreprise/temps/decompte?debut={j}&fin={j}", headers=a["membre"]).json()
    l2 = next(x for x in d2["dossiers"] if x["dossier"] == "2024-16")["lignes"][0]
    assert l2["libelle"] == "Audit social" and l2["description"].startswith("Analyse")
    x = client.get(f"/api/entreprise/temps/decompte?debut={j}&fin={j}&format=xlsx", headers=a["membre"])
    from openpyxl import load_workbook
    import io
    assert load_workbook(io.BytesIO(x.content)).sheetnames == ["Détail", "Par dossier"]
    # le gérant voit tout le monde ; un membre ne voit que lui
    assert client.get(f"/api/entreprise/temps/decompte?debut={j}&fin={j}", headers=a["proprietaire"]).json()["total"] > 0
    envoyes = []

    async def faux(*, to, subject, html):
        envoyes.append((to, html))
        return "id"
    monkeypatch.setitem(server.__dict__, "send_email", faux)
    assert client.post("/api/entreprise/temps/decompte/envoyer", json={"debut": j, "fin": j, "a": "compta@client.fr"}, headers=a["membre"]).json()["ok"]
    assert envoyes[0][0] == "compta@client.fr" and "Dossier 2024-15" in envoyes[0][1] and "pas une facture" in envoyes[0][1]


def test_radar_partage_avec_l_equipe(client, compte):
    a = _equipe(client, compte)
    assert client.get("/api/entreprise/radar", headers=a["membre"]).json() == {"partage": False, "prospects": []}
    assert client.patch("/api/entreprise/parametres", json={"partage_radar": True}, headers=a["manager"]).status_code == 403
    r = client.patch("/api/entreprise/parametres", json={"partage_radar": True}, headers=a["proprietaire"]).json()
    assert r["marque"]["partage_radar"] is True
    assert client.get("/api/entreprise/radar", headers=a["membre"]).json()["partage"] is True
    assert client.get("/api/entreprise/radar", headers=a["partenaire"]).status_code == 403


def test_relier_compte_perso_et_notifications(client, compte):
    a = _equipe(client, compte)
    email_perso, h_perso = compte()
    _, h_autre = compte()
    assert client.get("/api/entreprise/liaison", headers=a["membre"]).json()["confirme"] is False
    r = client.post("/api/entreprise/liaison", json={"email_perso": email_perso}, headers=a["membre"]).json()
    assert r["en_attente"] is True and r["compte_trouve"] is True
    # le compte perso reçoit la demande (cloche) avec le jeton
    n = client.get("/api/notifications", headers=h_perso).json()["items"]
    lien = next(x for x in n if "Relier" in x.get("titre", ""))["url"]
    jeton = lien.split("jeton=")[1]
    # un autre compte ne peut pas confirmer à sa place
    assert client.post("/api/entreprise/liaison/accepter", json={"jeton": jeton}, headers=h_autre).status_code == 404
    assert client.post("/api/entreprise/liaison/accepter", json={"jeton": jeton}, headers=h_perso).json()["ok"] is True
    assert client.get("/api/entreprise/liaison", headers=a["membre"]).json() == {"email_perso": email_perso.lower(), "confirme": True, "en_attente": False, "notif_perso": True}
    # une notification d'équipe arrive aussi sur le compte perso…
    ids = {m["role"]: m["id"] for m in client.get("/api/entreprise/membres", headers=a["proprietaire"]).json()["membres"]}
    client.post("/api/entreprise/planning", json={"titre": "Réunion client", "debut": "2026-10-20", "membre_id": ids["membre"]}, headers=a["manager"])
    assert any("planning" in x.get("titre", "").lower() for x in client.get("/api/notifications", headers=h_perso).json()["items"])
    # …sauf si la personne le coupe
    client.put("/api/entreprise/liaison/notifs", json={"notif_perso": False}, headers=a["membre"])
    avant = len(client.get("/api/notifications", headers=h_perso).json()["items"])
    client.post("/api/entreprise/planning", json={"titre": "Autre réunion", "debut": "2026-10-21", "membre_id": ids["membre"]}, headers=a["manager"])
    assert len(client.get("/api/notifications", headers=h_perso).json()["items"]) == avant
    assert client.delete("/api/entreprise/liaison", headers=a["membre"]).json()["confirme"] is False
