"""Installation en un clic : dossiers, classeurs, Documents-RH par personne, partages, relance sans doublon (faux Drive)."""
import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import server  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from tests.test_entreprise import _equipe  # noqa: E402


class FauxDrive:
    journal = []

    def __init__(self, c, jeton):
        self.n = 0

    async def dossier(self, nom, parent):
        FauxDrive.journal.append(("dossier", nom, parent))
        return {"id": f"d-{nom}-{len(FauxDrive.journal)}", "url": f"https://drive.google.com/drive/folders/{len(FauxDrive.journal)}"}

    async def classeur(self, nom, octets, parent):
        FauxDrive.journal.append(("classeur", nom, octets))
        return {"id": f"c-{nom}", "url": f"https://docs.google.com/spreadsheets/d/{nom}"}

    async def partager(self, item, email, ecrire):
        FauxDrive.journal.append(("partage", item, email, ecrire))


def test_installation_un_clic(client, compte, monkeypatch):
    a = _equipe(client, compte)
    # sans Drive relié : message clair
    assert client.post("/api/entreprise/installer", json={"fournisseur": "google"}, headers=a["proprietaire"]).status_code == 409
    assert client.post("/api/entreprise/installer", json={"fournisseur": "google"}, headers=a["manager"]).status_code == 403

    async def jeton(db, provider):
        return "jeton-test", None
    monkeypatch.setitem(server.__dict__, "_cloud_token", jeton)
    monkeypatch.setitem(server.__dict__, "_installateur_test", FauxDrive)
    FauxDrive.journal = []
    r = client.post("/api/entreprise/installer", json={"fournisseur": "google"}, headers=a["proprietaire"])
    assert r.status_code == 200, r.text
    b = r.json()
    noms = [x[1] for x in FauxDrive.journal if x[0] in ("dossier", "classeur")]
    assert noms[0].startswith("Zayado RH – ") and "Zayado RH – Référence" in noms and "Zayado RH – Données privées" in noms and "Documents-RH" in noms
    assert b["dossiers_crees"] >= 2 and noms.count("Bulletins") == b["dossiers_crees"]
    # l'Annuaire du classeur de référence est pré-rempli avec l'équipe
    ref = next(x[2] for x in FauxDrive.journal if x[0] == "classeur" and x[1] == "Zayado RH – Référence")
    lignes = list(load_workbook(io.BytesIO(ref))["Annuaire"].iter_rows(min_row=2, values_only=True))
    assert len(lignes) >= 3 and any(l[2] == "Employeur" for l in lignes)
    # chaque personne : son dossier partagé en écriture ; le classeur de référence en lecture (écriture pour les gérants)
    partages = [x for x in FauxDrive.journal if x[0] == "partage"]
    assert any(x[3] is True and x[1].startswith("d-") for x in partages) and any(x[1] == "c-Zayado RH – Référence" and x[3] is False for x in partages)
    # son bouton « Mon Drive pro » ouvre maintenant SON dossier
    assert client.get("/api/entreprise/mon-drive", headers=a["membre"]).json()["source"] == "personnel"
    etat = client.get("/api/entreprise/installation", headers=a["proprietaire"]).json()
    assert etat["installe"] is True and etat["fournisseur"] == "google"
    # relance : rien n'est recréé
    FauxDrive.journal = []
    b2 = client.post("/api/entreprise/installer", json={"fournisseur": "google"}, headers=a["proprietaire"]).json()
    assert b2["dossiers_crees"] == 0 and b2["partages"] == 0 and FauxDrive.journal == []


def test_requetes_google_et_microsoft():
    import asyncio
    import json as _json
    import httpx
    from entreprise_installation_ext import Google, Microsoft
    vues = []

    class C:  # faux client : enregistre chaque requête (le réseau est coupé pendant les tests)
        async def post(self, url, headers=None, json=None, files=None):
            corps = _json.dumps(json).encode() if json is not None else (files["metadata"][1].encode() if files else b"")
            vues.append(("POST", url, corps))
            return httpx.Response(200, json={"id": "X1", "webViewLink": "https://drive/x", "webUrl": "https://onedrive/x"})

        async def put(self, url, headers=None, content=None):
            vues.append(("PUT", url, content))
            return httpx.Response(200, json={"id": "X2", "webUrl": "https://onedrive/y"})

    async def go():
        gd = Google(C(), "jt")
        await gd.dossier("Zayado RH", None)
        await gd.classeur("Référence", b"xlsx", "P1")
        await gd.partager("X1", "lea@x.fr", False)
        ms = Microsoft(C(), "jt")
        await ms.dossier("Documents-RH", "P1")
        await ms.classeur("Référence", b"xlsx", "P1")
        await ms.partager("X1", "lea@x.fr", True)
    asyncio.run(go())
    assert vues[0][0] == "POST" and "drive/v3/files" in vues[0][1] and b"application/vnd.google-apps.folder" in vues[0][2]
    assert "upload/drive/v3/files?uploadType=multipart" in vues[1][1] and b"application/vnd.google-apps.spreadsheet" in vues[1][2]  # converti en Google Sheets
    assert "/files/X1/permissions" in vues[2][1] and _json.loads(vues[2][2]) == {"type": "user", "role": "reader", "emailAddress": "lea@x.fr"}
    assert vues[3][1].endswith("/me/drive/items/P1/children") and b'"folder"' in vues[3][2]
    assert vues[4][0] == "PUT" and "/me/drive/items/P1:/R%C3%A9f%C3%A9rence.xlsx:/content" in vues[4][1]
    assert vues[5][1].endswith("/items/X1/invite") and _json.loads(vues[5][2])["roles"] == ["write"] and _json.loads(vues[5][2])["requireSignIn"] is True
