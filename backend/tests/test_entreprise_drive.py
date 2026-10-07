"""« Ton entreprise » (lot du 5 oct., suite) : pièces demandées → déposées dans le Drive → validées ; planning indiqué par le
salarié → notification + commentaire du gérant ; coordonnées ; marque du client ; Excel (modèle, import, export)."""
import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import server  # noqa: E402
from entreprise_excel import classeur_modele, lire_classeur  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from tests.test_entreprise import _equipe, _ids  # noqa: E402


def _notifs(client, h):
    r = client.get("/api/notifications", headers=h)
    if r.status_code != 200:
        return []
    d = r.json()
    return d.get("items") or d.get("notifications") or (d if isinstance(d, list) else [])


def test_pieces_liste_drive_validation(client, compte):
    a = _equipe(client, compte)
    ids = _ids(client, a["proprietaire"])
    # le dirigeant dit où sont les dossiers
    r = client.patch("/api/entreprise/parametres", json={"drive_fournisseur": "microsoft", "drive_url": "https://acme.sharepoint.com/sites/RH",
                                                        "consignes": "Un fichier PDF par pièce."}, headers=a["proprietaire"])
    assert r.status_code == 200 and r.json()["marque"]["drive_url"].startswith("https://")
    assert client.patch("/api/entreprise/parametres", json={"drive_url": "javascript:alert(1)"}, headers=a["proprietaire"]).status_code == 422
    # sans dossier personnel : le bouton ouvre le dossier commun
    assert client.get("/api/entreprise/mon-drive", headers=a["membre"]).json()["source"] == "entreprise"
    # dossier personnel posé par le gérant, pas par le salarié
    perso = "https://acme.sharepoint.com/sites/RH/Documents-RH/Membre"
    client.put("/api/entreprise/fiches/moi", json={"lien_drive": "https://pirate.example/x"}, headers=a["membre"])
    assert client.get("/api/entreprise/mon-drive", headers=a["membre"]).json()["source"] == "entreprise"
    assert client.put(f"/api/entreprise/fiches/{ids['membre']}", json={"lien_drive": perso}, headers=a["manager"]).status_code == 200
    d = client.get("/api/entreprise/mon-drive", headers=a["membre"]).json()
    assert d["url"] == perso and d["source"] == "personnel" and d["consignes"]
    # liste de l'entreprise, demandée à toute l'équipe
    t = client.post("/api/entreprise/types-pieces", json={"titre": "RIB", "consigne": "Au nom de la personne"}, headers=a["proprietaire"]).json()
    assert client.post("/api/entreprise/types-pieces", json={"titre": "rib"}, headers=a["proprietaire"]).status_code == 409
    assert client.post("/api/entreprise/types-pieces", json={"titre": "X"}, headers=a["membre"]).status_code == 403
    r = client.post(f"/api/entreprise/types-pieces/{t['id']}/demander", json={}, headers=a["proprietaire"]).json()
    assert r["crees"] >= 2  # membre + manager (partenaire et propriétaire exclus)
    assert client.post(f"/api/entreprise/types-pieces/{t['id']}/demander", json={}, headers=a["proprietaire"]).json()["crees"] == 0
    mes = client.get("/api/entreprise/pieces", headers=a["membre"]).json()["pieces"]
    rib = next(p for p in mes if p["titre"] == "RIB")
    assert rib["statut"] == "a_fournir" and rib["consigne"] == "Au nom de la personne"
    # déposé (sans lien : il est dans le Drive) → le gérant voit « à vérifier »
    assert client.patch(f"/api/entreprise/pieces/{rib['id']}", json={"statut": "fournie"}, headers=a["membre"]).status_code == 200
    assert client.get("/api/entreprise/apercu", headers=a["manager"]).json()["pieces_a_verifier"] >= 1
    assert any("déposé" in n.get("titre", "") for n in _notifs(client, a["manager"]))
    assert client.patch(f"/api/entreprise/pieces/{rib['id']}", json={"statut": "validee"}, headers=a["membre"]).status_code == 403
    # refus motivé, nouveau dépôt, validation
    client.patch(f"/api/entreprise/pieces/{rib['id']}", json={"statut": "refusee", "commentaire": "Illisible"}, headers=a["manager"])
    p = next(x for x in client.get("/api/entreprise/pieces", headers=a["membre"]).json()["pieces"] if x["id"] == rib["id"])
    assert p["statut"] == "refusee" and p["commentaire"] == "Illisible"
    client.patch(f"/api/entreprise/pieces/{rib['id']}", json={"statut": "fournie"}, headers=a["membre"])
    p = next(x for x in client.get("/api/entreprise/pieces", headers=a["membre"]).json()["pieces"] if x["id"] == rib["id"])
    assert p["statut"] == "fournie" and p["commentaire"] == ""
    assert client.patch(f"/api/entreprise/pieces/{rib['id']}", json={"statut": "validee"}, headers=a["proprietaire"]).status_code == 200
    assert client.patch(f"/api/entreprise/pieces/{rib['id']}", json={"statut": "fournie"}, headers=a["membre"]).status_code == 409


def test_piece_par_defaut_posee_a_l_arrivee(client, compte):
    a = _equipe(client, compte)
    client.post("/api/entreprise/types-pieces", json={"titre": "Carte vitale", "par_defaut": True}, headers=a["proprietaire"])
    client.post("/api/entreprise/types-pieces", json={"titre": "Permis", "par_defaut": False}, headers=a["proprietaire"])
    r = client.post("/api/entreprise/membres", json={"role": "membre", "nom": "Nouveau"}, headers=a["proprietaire"]).json()
    _, h = compte()
    assert client.post("/api/entreprise/rejoindre", json={"token": r["lien"].split("token=")[1]}, headers=h).status_code == 200
    titres = {p["titre"] for p in client.get("/api/entreprise/pieces", headers=h).json()["pieces"]}
    assert titres == {"Carte vitale"}


def test_planning_indique_par_le_salarie(client, compte):
    a = _equipe(client, compte)
    r = client.post("/api/entreprise/planning", json={"titre": "Prospection", "debut": "2026-10-12", "creneau": "matin",
                                                      "membre_id": "quelquun-dautre"}, headers=a["membre"])
    assert r.status_code == 200
    pid = r.json()["id"]
    pl = client.get("/api/entreprise/planning?debut=2026-10-12&fin=2026-10-18", headers=a["manager"]).json()["planning"]
    e = next(x for x in pl if x["id"] == pid)
    assert e["statut"] == "propose" and e["creneau"] == "matin" and e["membre_id"] == _ids(client, a["proprietaire"])["membre"]
    assert any(x["id"] == pid for x in client.get("/api/entreprise/planning/a-regarder", headers=a["manager"]).json()["planning"])
    # le gérant est prévenu (cloche), avec un lien direct vers le planning
    assert any("planning" in (n.get("titre", "") + n.get("url", "")) for n in _notifs(client, a["manager"]))
    assert client.get("/api/entreprise/planning/a-regarder", headers=a["membre"]).status_code == 403
    assert client.post("/api/entreprise/planning", json={"titre": "X", "debut": "2026-10-12"}, headers=a["partenaire"]).status_code == 403
    # le gérant commente : c'est vu, le salarié lit le commentaire
    assert client.post("/api/entreprise/commentaires", json={"cible_type": "planning", "cible_id": pid, "texte": "OK, garde l'après-midi pour les devis"},
                       headers=a["manager"]).status_code == 200
    com = client.get(f"/api/entreprise/commentaires?cible_type=planning&cible_id={pid}", headers=a["membre"]).json()["commentaires"]
    assert len(com) == 1 and "devis" in com[0]["texte"]
    assert any("Commentaire" in n.get("titre", "") for n in _notifs(client, a["membre"]))
    pl = client.get("/api/entreprise/planning?debut=2026-10-12&fin=2026-10-18", headers=a["membre"]).json()["planning"]
    assert next(x for x in pl if x["id"] == pid)["statut"] == "vu"
    assert client.get(f"/api/entreprise/commentaires?cible_type=planning&cible_id={pid}", headers=a["partenaire"]).status_code == 403
    # le salarié supprime son entrée ; pas celle d'un autre
    autre = client.post("/api/entreprise/planning", json={"titre": "Réunion", "debut": "2026-10-13"}, headers=a["manager"]).json()["id"]
    assert client.delete(f"/api/entreprise/planning/{autre}", headers=a["membre"]).status_code == 403
    assert client.delete(f"/api/entreprise/planning/{pid}", headers=a["membre"]).status_code == 200


def test_coordonnees_privees_et_partage(client, compte):
    a = _equipe(client, compte)
    ids = _ids(client, a["proprietaire"])
    r = client.put("/api/entreprise/fiches/moi", json={"adresse": "3 rue des Lilas", "ville": "Lyon", "tel_perso": "0600000000",
                                                      "partage_tel": True, "tel_pro": "0102030405"}, headers=a["membre"])
    assert r.status_code == 200 and r.json()["tel_pro"] == ""  # l'info pro est décidée par l'entreprise
    assert client.get(f"/api/entreprise/fiches/{ids['membre']}", headers=a["partenaire"]).status_code == 403
    assert client.get(f"/api/entreprise/fiches/{ids['membre']}", headers=a["manager"]).json()["adresse"] == "3 rue des Lilas"
    # un gérant ne réécrit pas l'adresse à la place de la personne
    client.put(f"/api/entreprise/fiches/{ids['membre']}", json={"adresse": "ailleurs", "tel_pro": "0102030405"}, headers=a["manager"])
    f = client.get("/api/entreprise/fiches/moi", headers=a["membre"]).json()
    assert f["adresse"] == "3 rue des Lilas" and f["tel_pro"] == "0102030405"
    ann = {x["id"]: x for x in client.get("/api/entreprise/annuaire", headers=a["partenaire"]).json()["annuaire"]}
    assert ann[ids["membre"]]["tel"] == "0600000000" and ann[ids["membre"]]["email"] == "" and "adresse" not in ann[ids["membre"]]
    assert any(x["action"] == "fiche_lue" for x in client.get("/api/entreprise/journal", headers=a["proprietaire"]).json()["journal"])
    ex = client.get(f"/api/entreprise/membres/{ids['membre']}/export", headers=a["membre"]).json()
    assert ex["coordonnees"]["ville"] == "Lyon"


def test_marque_du_client(client, compte):
    a = _equipe(client, compte)
    logo = "data:image/png;base64,iVBORw0KGgo="
    r = client.patch("/api/entreprise/parametres", json={"nom_affiche": "Atelier Lumière", "logo": logo, "couleur": "#2f6f5e"}, headers=a["proprietaire"])
    assert r.status_code == 200
    m = client.get("/api/entreprise/moi", headers=a["membre"]).json()["marque"]
    assert m["nom"] == "Atelier Lumière" and m["logo"] == logo and m["couleur"] == "#2f6f5e"
    assert client.patch("/api/entreprise/parametres", json={"logo": "data:text/html;base64,PHNjcmlwdD4="}, headers=a["proprietaire"]).status_code == 422
    assert client.patch("/api/entreprise/parametres", json={"nom_affiche": "X"}, headers=a["manager"]).status_code == 403


def test_excel_modele_import_export(client, compte):
    a = _equipe(client, compte)
    assert client.get("/api/entreprise/excel/modele", headers=a["membre"]).status_code == 403
    r = client.get("/api/entreprise/excel/modele", headers=a["proprietaire"])
    assert r.status_code == 200 and {k for k, v in lire_classeur(r.content).items() if v} == {"TypesPieces"}  # le modèle se relit ; les lignes « Léa Martin » sont ignorées
    # un fichier rempli par l'entreprise
    wb = load_workbook(io.BytesIO(classeur_modele()))
    wb["Annuaire"].append(["Paul Durand", "paul@acme.fr", "Salarie", "Comptable", "0611", ""])
    wb["TypesPieces"].append(["Diplôme", "Non", ""])
    wb["PiecesRequises"].append(["Diplôme", "paul@acme.fr"])
    wb["Planning"].append(["Clôture", "Paul Durand", "2026-10-30", "Après-midi", "Bilan", ""])
    wb["Planning"].append(["Mauvaise date", "Paul Durand", "30/10", "", "", ""])
    wb["DossierSalarie"].append(["Paul Durand", "paul@acme.fr", "https://acme.sharepoint.com/RH/Paul"])
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post("/api/entreprise/excel/import", files={"fichier": ("equipe.xlsx", buf.getvalue(), "application/octet-stream")}, headers=a["proprietaire"])
    assert r.status_code == 200, r.text
    rap = r.json()
    assert rap["invites"] == 1 and rap["types"] == 2 and rap["pieces"] == 1 and rap["planning"] == 1 and rap["ignores"]
    # ré-import : personne n'est dupliqué
    r2 = client.post("/api/entreprise/excel/import", files={"fichier": ("equipe.xlsx", buf.getvalue(), "application/octet-stream")}, headers=a["proprietaire"]).json()
    assert r2["invites"] == 0 and r2["mis_a_jour"] == 1 and r2["pieces"] == 0
    membres = client.get("/api/entreprise/membres", headers=a["proprietaire"]).json()["membres"]
    assert sum(1 for m in membres if m.get("email") == "paul@acme.fr") == 1
    assert client.post("/api/entreprise/excel/import", files={"fichier": ("x.xlsx", b"pas un excel", "application/octet-stream")},
                       headers=a["proprietaire"]).status_code == 422
    # export : coordonnées personnelles pour le propriétaire seulement
    client.put("/api/entreprise/fiches/moi", json={"adresse": "3 rue des Lilas"}, headers=a["membre"])
    def colonnes(h):
        wb = load_workbook(io.BytesIO(client.get("/api/entreprise/excel/export", headers=h).content))
        ws = wb["DossierSalarie"]
        return [c.value for c in ws[1]], [list(r) for r in ws.iter_rows(min_row=2, values_only=True)]
    cols, lignes = colonnes(a["proprietaire"])
    assert "Adresse" in cols and any("3 rue des Lilas" in (l or []) for l in lignes)
    cols, _ = colonnes(a["manager"])
    assert "Adresse" not in cols


def test_kit_microsoft_reconnu():
    """Les fichiers unitaires du kit Microsoft Lists s'importent tels quels ; Pieces/Absences ne sont pas confondus."""
    kit = Path(__file__).resolve().parents[2] / "frontend" / "public" / "kits"
    if not kit.exists():
        return
    import zipfile
    z = kit / "Kit-installation-ZayadoRH.zip"
    if not z.exists():
        return
    with zipfile.ZipFile(z) as zz, zipfile.ZipFile(io.BytesIO(zz.read("Listes-import.zip"))) as listes:
        assert set(lire_classeur(listes.read("Planning.xlsx"))) == {"Planning"}
        try:
            lire_classeur(listes.read("Pieces.xlsx"))
            assert False, "Pieces.xlsx ne doit pas être lu comme un planning"
        except ValueError:
            pass
