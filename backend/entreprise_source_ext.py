"""« J'ai déjà mes fichiers » : le dirigeant relie SON classeur existant (Excel dans OneDrive/SharePoint, ou Google Sheets)
au lieu de laisser Zayado tout créer.

1. Il choisit son fichier (liste de ses classeurs Excel du OneDrive, ou classeurs Google auxquels Zayado a accès).
2. Zayado lit les onglets et leurs colonnes, et PROPOSE la correspondance (« ta colonne « Nom » = Nom de la personne ? »).
3. Il corrige si besoin et valide. 4. « Synchroniser » relit son fichier et met Zayado à jour (personnes, pièces demandées,
planning) — son fichier reste la référence ; rien n'est supprimé chez lui ni dans Zayado.

Droits : ceux que Zayado demande déjà (Microsoft : fichiers du OneDrive ; Google : fichiers créés par Zayado ou ouverts avec
lui). Un classeur Google existant qui n'apparaît pas doit être ouvert une fois avec Zayado (sélecteur Google) ou copié.
"""
import io
import logging
import re
import unicodedata
from datetime import date, datetime, timezone
from typing import Optional

import httpx
from fastapi import Depends, HTTPException
from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

log = logging.getLogger("kairos.source")

# Ce que Zayado sait lire, et les noms de colonnes qu'on rencontre souvent chez les clients
CIBLES = {
    "personnes": {"titre": "Les personnes", "champs": {
        "nom": ("Nom de la personne", True, ["title", "nom", "nom prenom", "prenom nom", "name", "salarie", "collaborateur", "employe", "personne"]),
        "email": ("E-mail", True, ["email", "e-mail", "mail", "courriel", "adresse mail", "adresse e-mail"]),
        "role": ("Rôle (Employeur / Salarié)", False, ["role", "statut", "type"]),
        "poste": ("Poste", False, ["poste", "fonction", "job", "metier", "intitule"]),
        "tel_pro": ("Téléphone pro", False, ["telpro", "tel pro", "telephone", "tel", "portable", "mobile"]),
    }},
    "pieces": {"titre": "Les pièces demandées", "champs": {
        "titre": ("Nom de la pièce", True, ["title", "piece", "document", "type", "nom", "libelle"]),
        "par_defaut": ("Demandée à chaque arrivée (Oui/Non)", False, ["pardefaut", "par defaut", "obligatoire", "systematique"]),
        "consigne": ("Consigne", False, ["consigne", "indication", "commentaire", "description"]),
    }},
    "planning": {"titre": "Le planning", "champs": {
        "titre": ("Tâche / événement", True, ["title", "tache", "titre", "activite", "intitule", "mission", "objet"]),
        "date": ("Date", True, ["dateecheance", "date", "jour", "echeance", "le"]),
        "personne": ("Personne (nom ou e-mail)", False, ["salarie", "personne", "collaborateur", "responsable", "qui", "assigne"]),
        "creneau": ("Créneau (Matin / Après-midi)", False, ["creneau", "moment", "periode"]),
        "note": ("Précision / livrable", False, ["livrable", "commentaire", "note", "description", "detail"]),
    }},
}
# Onglet du client qui ressemble le plus à chaque cible (noms des listes du kit compris)
NOMS_ONGLETS = {"personnes": ["annuaire", "personnes", "salaries", "equipe", "collaborateurs", "employes", "rh"],
                "pieces": ["typespieces", "pieces", "documents", "pieces demandees"],
                "planning": ["planning", "agenda", "taches", "calendrier"]}


def norm(t) -> str:
    t = unicodedata.normalize("NFD", str(t or "").lower())
    t = "".join(c for c in t if unicodedata.category(c) != "Mn")
    return re.sub(r"[^a-z0-9]+", " ", t).strip()


def lire_xlsx(octets: bytes, n_exemples: int = 3) -> dict:
    """{onglet: {"colonnes": [...], "exemples": [[...]], "lignes": [[...]]}} (lignes plafonnées à 2 000)."""
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(octets), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001
        raise ValueError("Ce fichier n'est pas un classeur lisible.")
    out = {}
    for ws in wb.worksheets:
        it = ws.iter_rows(values_only=True)
        try:
            entete = [str(c).strip() if c is not None else "" for c in next(it)]
        except StopIteration:
            continue
        if not any(entete):
            continue
        lignes = []
        for n, r in enumerate(it):
            if n >= 2000:
                break
            if any(v not in (None, "") for v in r):
                lignes.append([v.isoformat()[:10] if isinstance(v, (datetime, date)) else ("" if v is None else str(v).strip()) for v in r])
        out[ws.title] = {"colonnes": entete, "exemples": lignes[:n_exemples], "lignes": lignes}
    if not out:
        raise ValueError("Aucun onglet avec des en-têtes n'a été trouvé.")
    return out


def proposer(onglets: dict) -> dict:
    """Pour chaque cible : l'onglet le plus probable et la colonne proposée pour chaque champ (ou rien si doute)."""
    prop = {}
    for cible, conf in CIBLES.items():
        meilleur, score = None, 0
        for nom, o in onglets.items():
            cols = {norm(c) for c in o["colonnes"]}
            s = sum(1 for _, (_, _, syn) in conf["champs"].items() if cols & {norm(x) for x in syn})
            s += 2 if norm(nom).replace(" ", "") in [x.replace(" ", "") for x in NOMS_ONGLETS[cible]] else 0
            if s > score:
                meilleur, score = nom, s
        if not meilleur or score < 2:
            prop[cible] = None
            continue
        cols = onglets[meilleur]["colonnes"]
        champs, pris = {}, set()
        for champ, (_, _, syn) in conf["champs"].items():
            choix = next((c for c in cols if c and c not in pris and norm(c) in {norm(x) for x in syn}), None)
            if choix is None:  # correspondance partielle (« Adresse e-mail pro » ⊃ « mail »)
                choix = next((c for c in cols if c and c not in pris and any(norm(x) and norm(x) in norm(c) for x in syn if len(norm(x)) > 3)), None)
            if choix:
                champs[champ] = choix
                pris.add(choix)
        prop[cible] = {"onglet": meilleur, "colonnes": champs}
    return prop


def convertir(onglets: dict, correspondance: dict) -> dict:
    """Les lignes du client → le format d'import de Zayado (noms de colonnes du kit)."""
    out = {}
    for cible, c in (correspondance or {}).items():
        if not c or c.get("onglet") not in onglets:
            continue
        o = onglets[c["onglet"]]
        idx = {champ: o["colonnes"].index(col) for champ, col in (c.get("colonnes") or {}).items() if col in o["colonnes"]}
        val = lambda l, ch: (l[idx[ch]] if ch in idx and idx[ch] < len(l) else "")  # noqa: E731
        if cible == "personnes":
            out["Annuaire"] = [{"Title": val(l, "nom"), "Email": val(l, "email"),
                                "Role": "Employeur" if norm(val(l, "role")) in ("employeur", "dirigeant", "manager", "gerant", "responsable") else "Salarie",
                                "Poste": val(l, "poste"), "TelPro": val(l, "tel_pro")} for l in o["lignes"]]
        elif cible == "pieces":
            out["TypesPieces"] = [{"Title": val(l, "titre"), "ParDefaut": "Non" if norm(val(l, "par_defaut")) in ("non", "no", "0", "faux") else "Oui",
                                   "Consigne": val(l, "consigne")} for l in o["lignes"]]
        elif cible == "planning":
            out["Planning"] = [{"Title": val(l, "titre"), "DateEcheance": val(l, "date")[:10], "Salarie": val(l, "personne"),
                                "Creneau": val(l, "creneau"), "Commentaire": val(l, "note")} for l in o["lignes"]]
    return out


def install_entreprise_source(g: dict) -> None:
    api, get_db = g["api"], g["get_db"]
    EntMembre, EntReglage = g["EntMembre"], g["EntReglage"]

    async def _dirigeant(db):
        uid = g["_uid"]()
        m = (await db.execute(select(EntMembre).where(EntMembre.user_id == uid, EntMembre.statut == "actif"))).scalars().first()
        if not m or m.role != "proprietaire":
            raise HTTPException(403, "Seul le dirigeant relie les fichiers de l'entreprise.")
        return m

    async def _jeton(db, fournisseur: str) -> str:
        try:
            jeton, _ = await g["_cloud_token"](db, fournisseur)
            return jeton
        except HTTPException:
            raise HTTPException(409, f"Relie d'abord ton {'Google Drive' if fournisseur == 'google' else 'OneDrive'} (Paramètres › Connexions).")

    async def chercher(db, fournisseur: str, q: str) -> list:
        f = g.get("_source_test_chercher")
        if f:
            return await f(fournisseur, q)
        h = {"Authorization": f"Bearer {await _jeton(db, fournisseur)}"}
        async with httpx.AsyncClient(timeout=20) as c:
            if fournisseur == "microsoft":
                r = await c.get(f"https://graph.microsoft.com/v1.0/me/drive/root/search(q='{(q or 'xlsx').replace(chr(39), '')}')", headers=h,
                                params={"$select": "id,name,webUrl,file", "$top": "50"})
                items = r.json().get("value", []) if r.status_code < 400 else []
                return [{"id": i["id"], "nom": i["name"], "url": i.get("webUrl")} for i in items if i.get("name", "").lower().endswith((".xlsx", ".xlsm"))]
            requete = "mimeType='application/vnd.google-apps.spreadsheet' and trashed=false" + (f" and name contains '{q.replace(chr(39), '')}'" if q else "")
            r = await c.get("https://www.googleapis.com/drive/v3/files", headers=h, params={"q": requete, "fields": "files(id,name,webViewLink)", "pageSize": "50"})
            return [{"id": i["id"], "nom": i["name"], "url": i.get("webViewLink")} for i in (r.json().get("files", []) if r.status_code < 400 else [])]

    async def telecharger(db, fournisseur: str, fichier_id: str) -> bytes:
        f = g.get("_source_test_telecharger")
        if f:
            return await f(fournisseur, fichier_id)
        h = {"Authorization": f"Bearer {await _jeton(db, fournisseur)}"}
        async with httpx.AsyncClient(timeout=60, follow_redirects=True) as c:
            if fournisseur == "microsoft":
                r = await c.get(f"https://graph.microsoft.com/v1.0/me/drive/items/{fichier_id}/content", headers=h)
            else:
                r = await c.get(f"https://www.googleapis.com/drive/v3/files/{fichier_id}/export", headers=h,
                                params={"mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"})
        if r.status_code >= 400:
            raise HTTPException(502, "Impossible d'ouvrir ce fichier avec les droits accordés à Zayado.")
        return r.content

    class AnalyserIn(BaseModel):
        fournisseur: str
        fichier_id: str = Field(min_length=1, max_length=300)

    class SourceIn(BaseModel):
        fournisseur: str
        fichier: dict
        correspondance: dict

    @api.get("/entreprise/source/fichiers")
    async def source_fichiers(fournisseur: str, q: str = "", db: AsyncSession = Depends(get_db)):
        await _dirigeant(db)
        if fournisseur not in ("google", "microsoft"):
            raise HTTPException(422, "Choisis Google ou Microsoft.")
        return {"fichiers": await chercher(db, fournisseur, q.strip()[:80])}

    @api.post("/entreprise/source/analyser")
    async def source_analyser(body: AnalyserIn, db: AsyncSession = Depends(get_db)):
        await _dirigeant(db)
        try:
            onglets = lire_xlsx(await telecharger(db, body.fournisseur, body.fichier_id))
        except ValueError as e:
            raise HTTPException(422, str(e))
        return {"onglets": {k: {"colonnes": v["colonnes"], "exemples": v["exemples"], "lignes": len(v["lignes"])} for k, v in onglets.items()},
                "proposition": proposer(onglets),
                "cibles": {k: {"titre": c["titre"], "champs": {ch: {"libelle": l, "obligatoire": o} for ch, (l, o, _) in c["champs"].items()}} for k, c in CIBLES.items()}}

    @api.put("/entreprise/source")
    async def source_enregistrer(body: SourceIn, db: AsyncSession = Depends(get_db)):
        m = await _dirigeant(db)
        for cible, c in body.correspondance.items():
            if cible not in CIBLES:
                raise HTTPException(422, f"Cible inconnue : {cible}.")
            if c:
                manque = [CIBLES[cible]["champs"][ch][0] for ch, (_, oblig, _) in CIBLES[cible]["champs"].items() if oblig and not (c.get("colonnes") or {}).get(ch)]
                if manque:
                    raise HTTPException(422, f"{CIBLES[cible]['titre']} : choisis la colonne pour « {', '.join(manque)} ».")
        reg = await db.get(EntReglage, m.org_id) or EntReglage(org_id=m.org_id, secu_actif=False)
        inst = dict(reg.installation or {})
        inst["source"] = {"fournisseur": body.fournisseur, "fichier": {k: body.fichier.get(k) for k in ("id", "nom", "url")},
                          "correspondance": body.correspondance}
        reg.installation = inst
        db.add(reg)
        await db.commit()
        return {"ok": True}

    @api.post("/entreprise/source/synchroniser")
    async def source_synchroniser(db: AsyncSession = Depends(get_db)):
        m = await _dirigeant(db)
        reg = await db.get(EntReglage, m.org_id)
        src = ((reg.installation if reg else None) or {}).get("source")
        if not src:
            raise HTTPException(409, "Aucun fichier relié.")
        onglets = lire_xlsx(await telecharger(db, src["fournisseur"], src["fichier"]["id"]))
        rapport = await g["_importer_entreprise"](db, m, convertir(onglets, src["correspondance"]))
        inst = dict(reg.installation or {})
        inst["source"] = {**src, "derniere_synchro": datetime.now(timezone.utc).isoformat()}
        reg.installation = inst
        await db.commit()
        return rapport

    @api.get("/entreprise/source")
    async def source_etat(db: AsyncSession = Depends(get_db)):
        m = await _dirigeant(db)
        reg = await db.get(EntReglage, m.org_id)
        return {"source": ((reg.installation if reg else None) or {}).get("source")}

    @api.delete("/entreprise/source")
    async def source_delier(db: AsyncSession = Depends(get_db)):
        m = await _dirigeant(db)
        reg = await db.get(EntReglage, m.org_id)
        if reg and reg.installation and "source" in reg.installation:
            inst = dict(reg.installation)
            inst.pop("source")
            reg.installation = inst
            await db.commit()
        return {"ok": True}
