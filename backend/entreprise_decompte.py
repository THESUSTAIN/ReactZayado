"""Décompte mensuel de « Ton entreprise » : présence, absences (dont congés payés), temps école, temps saisi, planning,
et paie ESTIMÉE (total à payer / déjà versé / reste à verser). Fonctions pures, sans base : testées seules.

Ce n'est pas un logiciel de paie : c'est un relevé pour préparer la paie et pour que chacun voie où il en est.
Congés payés : jours ouvrés (lundi-vendredi), 2,08 jours acquis par mois travaillé par défaut (25 jours par an).
"""
import io
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

LIB_ABSENCE = {"conges": "Congés payés (CP)", "maladie": "Maladie", "recuperation": "Récupération",
               "ecole": "Temps école", "sans_solde": "Sans solde", "autre": "Autre"}
NON_PAYEES = ("maladie", "sans_solde", "autre")   # retirées du dû ; la maladie est indemnisée à part (IJSS, maintien de salaire)
LIB_PRESENCE = {"bureau": "Bureau", "teletravail": "Télétravail", "deplacement": "Déplacement", "absent": "Absent"}


def bornes_mois(mois: str) -> tuple:
    a, m = (int(x) for x in mois.split("-"))
    debut = date(a, m, 1)
    fin = date(a + (m == 12), m % 12 + 1, 1) - timedelta(days=1)
    return debut, fin


def jours_ouvres(debut: date, fin: date) -> list:
    out, j = [], debut
    while j <= fin:
        if j.weekday() < 5:
            out.append(j.isoformat())
        j += timedelta(days=1)
    return out


def mois_entre(depuis: str, fin: date) -> int:
    """Mois entamés depuis `depuis` (inclus) jusqu'au mois de `fin` (inclus)."""
    d = date.fromisoformat(depuis)
    return max(0, (fin.year - d.year) * 12 + fin.month - d.month + 1)


def calculer(mois: str, presences: list, absences: list, temps: list, planning: list, remu: dict, cp_par_mois: float,
             absences_cp_depuis: list, verse: float, montrer_paie: bool) -> dict:
    """presences : [{jour, statut}] du mois ; absences : acceptées [{type, debut, fin}] (toutes dates) ;
    temps : [{jour, heures}] du mois ; planning : [{titre, debut, fin, creneau, statut}] du mois ;
    remu : {mode, taux, heures_jour, cp_initial, cp_depuis, ecole_jours: [0..4], ecole_payee} ;
    absences_cp_depuis : congés payés acceptés [{debut, fin}] depuis le début du compteur."""
    debut, fin = bornes_mois(mois)
    ouvres = jours_ouvres(debut, fin)
    ens = set(ouvres)
    # Jours d'absence acceptée, par type, dans le mois (un jour = un seul type : le premier posé l'emporte)
    jour_type = {}
    for a in absences:
        d, f = date.fromisoformat(a["debut"]), date.fromisoformat(a["fin"])
        for j in jours_ouvres(max(d, debut), min(f, fin)):
            jour_type.setdefault(j, a["type"])
    # Temps école récurrent (jours de la semaine choisis par l'admin), s'il n'y a pas déjà une absence ce jour-là
    for j in ouvres:
        if date.fromisoformat(j).weekday() in (remu.get("ecole_jours") or []) and j not in jour_type:
            jour_type[j] = "ecole"
    par_type = {k: 0 for k in LIB_ABSENCE}
    for j, t in jour_type.items():
        if j in ens:
            par_type[t if t in par_type else "autre"] += 1
    pres = {k: 0 for k in LIB_PRESENCE}
    for p in presences:
        if p["jour"] in ens and p["statut"] in pres:
            pres[p["statut"]] += 1
    heures_saisies = round(sum(float(t["heures"]) for t in temps if debut.isoformat() <= t["jour"] <= fin.isoformat()), 2)

    # Congés payés : acquis depuis le départ du compteur, pris (acceptés) jusqu'à la fin du mois, solde
    cp = None
    if remu.get("cp_depuis"):
        acquis = round(float(remu.get("cp_initial") or 0) + cp_par_mois * mois_entre(remu["cp_depuis"], fin), 2)
        pris = 0
        for a in absences_cp_depuis:
            d, f = date.fromisoformat(a["debut"]), date.fromisoformat(a["fin"])
            pris += len(jours_ouvres(max(d, date.fromisoformat(remu["cp_depuis"])), min(f, fin)))
        cp = {"acquis": acquis, "pris": pris, "pris_ce_mois": par_type["conges"], "solde": round(acquis - pris, 2)}

    absents = sum(par_type.values())
    travailles = len(ouvres) - absents
    out = {"mois": mois, "jours_ouvres": len(ouvres), "jours_travailles": travailles, "presence": pres, "absences": par_type,
           "cp": cp, "heures_saisies": heures_saisies, "planning": planning,
           "planning_nb": len(planning), "paie": None}
    if montrer_paie and remu.get("mode") in ("horaire", "forfait_jour") and float(remu.get("taux") or 0) > 0:
        hj = float(remu.get("heures_jour") or 7)
        payes = travailles + par_type["conges"] + par_type["recuperation"] + (par_type["ecole"] if remu.get("ecole_payee") else 0)
        if remu["mode"] == "horaire":
            # Heures réellement saisies si la personne saisit son temps, sinon jours payés × heures par jour
            base_heures = heures_saisies + (par_type["conges"] + par_type["recuperation"] + (par_type["ecole"] if remu.get("ecole_payee") else 0)) * hj \
                if heures_saisies > 0 else payes * hj
            quantite, unite = round(base_heures, 2), "h"
        else:
            quantite, unite = payes, "j"
        total = round(quantite * float(remu["taux"]), 2)
        out["paie"] = {"mode": remu["mode"], "quantite": quantite, "unite": unite, "taux": float(remu["taux"]),
                       "total_a_payer": total, "deja_verse": round(float(verse or 0), 2), "reste_a_verser": round(total - float(verse or 0), 2)}
    return out


_ENTETE = Font(bold=True, color="FFFFFF")
_FOND = PatternFill("solid", fgColor="1F3A5F")


def _entete(ws, cols):
    ws.append(cols)
    for c in ws[ws.max_row]:
        c.font, c.fill = _ENTETE, _FOND


def classeur(org: str, mois: str, fiches: list) -> bytes:
    """fiches : [(nom, decompte)] → un onglet Synthèse + un onglet par personne."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Synthèse"
    ws.append([f"{org} : compte rendu de {mois}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["Relevé préparatoire : ne remplace pas un bulletin de paie."])
    ws.append([])
    _entete(ws, ["Personne", "Jours ouvrés", "Jours travaillés", "CP pris", "Solde CP", "Maladie", "École", "Autres absences",
                 "Heures saisies", "Total à payer (€)", "Déjà versé (€)", "Reste à verser (€)"])
    for nom, d in fiches:
        a, p, cp = d["absences"], d["paie"] or {}, d["cp"] or {}
        ws.append([nom, d["jours_ouvres"], d["jours_travailles"], a["conges"], cp.get("solde", ""), a["maladie"], a["ecole"],
                   a["recuperation"] + a["sans_solde"] + a["autre"], d["heures_saisies"],
                   p.get("total_a_payer", ""), p.get("deja_verse", ""), p.get("reste_a_verser", "")])
    ws.column_dimensions["A"].width = 26
    for nom, d in fiches:
        w = wb.create_sheet((nom or "Personne")[:28].replace("/", "-"))
        w.append([f"{nom} · {mois}"])
        w["A1"].font = Font(bold=True, size=13)
        w.append([])
        _entete(w, ["Jours", "Nombre"])
        w.append(["Jours ouvrés du mois", d["jours_ouvres"]])
        w.append(["Jours travaillés", d["jours_travailles"]])
        for k, v in d["presence"].items():
            w.append([f"Présence déclarée : {LIB_PRESENCE[k]}", v])
        w.append([])
        _entete(w, ["Absences du mois", "Jours"])
        for k, v in d["absences"].items():
            w.append([LIB_ABSENCE[k], v])
        if d["cp"]:
            w.append([])
            _entete(w, ["Congés payés (CP)", "Jours"])
            for k, l in (("acquis", "Acquis au total"), ("pris", "Pris au total"), ("pris_ce_mois", "Pris ce mois"), ("solde", "Solde")):
                w.append([l, d["cp"][k]])
        w.append([])
        w.append(["Heures saisies", d["heures_saisies"]])
        if d["paie"]:
            p = d["paie"]
            w.append([])
            _entete(w, ["Paie (estimation)", "Valeur"])
            w.append([f"Base ({'heures' if p['unite'] == 'h' else 'jours'})", p["quantite"]])
            w.append([f"Taux ({'€ / heure' if p['unite'] == 'h' else '€ / jour'})", p["taux"]])
            w.append(["Total à payer (€)", p["total_a_payer"]])
            w.append(["Déjà versé (€)", p["deja_verse"]])
            w.append(["Reste à verser (€)", p["reste_a_verser"]])
        if d["planning"]:
            w.append([])
            _entete(w, ["Planning", "Du", "Au", "Créneau", "Statut"])
            for r in d["planning"]:
                w.append([r["titre"], r["debut"], r.get("fin") or "", r.get("creneau") or "", r.get("statut") or ""])
        w.column_dimensions["A"].width = 34
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
